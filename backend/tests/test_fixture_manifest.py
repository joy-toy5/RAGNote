import hashlib
import json
import zipfile
from pathlib import Path

import fitz
import pytest
from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from fixtures.generate_fixtures import build_assets


def _office_text(path: Path) -> str:
    if path.suffix == ".docx":
        document = WordDocument(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + cells)

    presentation = Presentation(path)
    values: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                values.append(shape.text)
            if getattr(shape, "has_table", False):
                values.extend(
                    cell.text for row in shape.table.rows for cell in row.cells
                )
    return "\n".join(values)


def test_fixture_manifest_has_two_isolated_users(backend_root: Path) -> None:
    fixture_root = backend_root / "tests/fixtures"
    manifest = json.loads((fixture_root / "manifest.json").read_text(encoding="utf-8"))

    users = {user["user_id"] for user in manifest["users"]}
    assert users == {"eval-user-a", "eval-user-b"}
    assert {document["owner"] for document in manifest["documents"]} == users

    contents: list[bytes] = []
    for document in manifest["documents"]:
        path = (fixture_root / document["path"]).resolve()
        assert path.is_relative_to(fixture_root.resolve())
        assert path.is_file()
        content = path.read_bytes()
        assert hashlib.sha256(content).hexdigest() == document["sha256"]
        assert document["synthetic"] is True
        contents.append(content)

    assert all(contents)
    assert len(set(contents)) == len(contents)
    for field in ("document_id", "path", "format", "sha256", "marker"):
        values = [document[field] for document in manifest["documents"]]
        assert len(values) == len(set(values)), f"{field} 必须唯一"


def test_fixture_manifest_covers_supported_ingestion_formats(
    backend_root: Path,
) -> None:
    fixture_root = backend_root / "tests/fixtures"
    manifest = json.loads((fixture_root / "manifest.json").read_text(encoding="utf-8"))
    by_format = {document["format"]: document for document in manifest["documents"]}

    assert set(by_format) == {
        "txt",
        "md",
        "pdf-text",
        "pdf-chart",
        "pdf-scan",
        "pdf-corrupt",
        "docx",
        "pptx",
    }

    for format_name in ("txt", "md"):
        item = by_format[format_name]
        content = (fixture_root / item["path"]).read_text(encoding="utf-8")
        assert item["marker"] in content

    for format_name in ("pdf-text", "pdf-chart", "pdf-scan"):
        item = by_format[format_name]
        with fitz.open(fixture_root / item["path"]) as document:
            assert document.page_count == item["page_count"]
            page = document[0]
            assert bool(page.get_images()) is item["has_embedded_images"]
            if format_name == "pdf-scan":
                assert page.get_text().strip() == ""
            else:
                assert item["marker"] in page.get_text()
            if item["has_embedded_images"]:
                image_xref = page.get_images(full=True)[0][0]
                image = document.extract_image(image_xref)["image"]
                assert hashlib.sha256(image).hexdigest() == item["image_sha256"]

    text_pdf = by_format["pdf-text"]
    with fitz.open(fixture_root / text_pdf["path"]) as document:
        assert (
            len(document[0].get_text().strip())
            >= text_pdf["minimum_native_text_length"]
        )
    assert text_pdf["requires_vision"] is False
    assert by_format["pdf-chart"]["requires_vision"] is True
    assert by_format["pdf-scan"]["requires_vision"] is True

    docx_item = by_format["docx"]
    docx_path = fixture_root / docx_item["path"]
    docx = WordDocument(docx_path)
    assert docx_item["marker"] in _office_text(docx_path)
    assert docx.core_properties.last_modified_by == "RAG-Note M0"
    assert len(docx.tables) == 1
    assert len(docx.tables[0].rows) == 2
    assert len(docx.tables[0].columns) == 2
    assert docx.tables[0].cell(1, 1).text == "bounded retry"

    pptx_item = by_format["pptx"]
    pptx_path = fixture_root / pptx_item["path"]
    presentation = Presentation(pptx_path)
    assert pptx_item["marker"] in _office_text(pptx_path)
    assert presentation.core_properties.last_modified_by == "RAG-Note M0"
    assert len(presentation.slides) == pptx_item["slide_count"]
    tables = [
        shape.table
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
    ]
    assert len(tables) == 1
    assert len(tables[0].rows) == 2
    assert len(tables[0].columns) == 2
    assert tables[0].cell(1, 1).text == "M0_PPTX_BOUNDED_RETRY"

    for path in (docx_path, pptx_path):
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None

    corrupt = by_format["pdf-corrupt"]
    corrupt_path = fixture_root / corrupt["path"]
    assert corrupt_path.read_bytes().startswith(corrupt["marker"].encode("ascii"))
    with fitz.open(corrupt_path) as document:
        assert document.page_count == 0
    with pytest.raises(PdfReadError, match="EOF marker not found"):
        PdfReader(corrupt_path, strict=True)


def test_generated_binary_fixtures_are_reproducible(backend_root: Path) -> None:
    corpus_root = backend_root / "tests/fixtures/corpus"
    for name, expected in build_assets().items():
        assert (corpus_root / name).read_bytes() == expected
