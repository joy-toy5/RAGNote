from __future__ import annotations

import argparse
import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import fitz
from docx import Document as WordDocument
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

FIXTURE_ROOT = Path(__file__).resolve().parent
CORPUS_ROOT = FIXTURE_ROOT / "corpus"
FIXED_TIME = datetime(2020, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
ZIP_TIME = (2020, 1, 2, 3, 4, 4)


def _png_bytes(kind: str) -> bytes:
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    if kind == "chart":
        draw.line((60, 300, 590, 300), fill="black", width=3)
        draw.line((60, 40, 60, 300), fill="black", width=3)
        for x, height, color in (
            (130, 100, "#d1495b"),
            (280, 170, "#00798c"),
            (430, 240, "#edae49"),
        ):
            draw.rectangle((x, 300 - height, x + 80, 300), fill=color)
        draw.text((70, 15), "M0_PDF_CHART_QPS_42", fill="black")
    else:
        draw.rectangle((30, 30, 610, 330), outline="black", width=3)
        draw.text((70, 145), "M0_PDF_SCAN_RETRY_POLICY", fill="black")
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=False)
    return output.getvalue()


def _pdf_bytes(kind: str) -> bytes:
    document = fitz.open()
    page = document.new_page(width=640, height=360)
    if kind == "text":
        text = (
            "M0_PDF_TEXT_RETRY_POLICY. Requests returning HTTP 429 use bounded "
            "backoff and never persist empty vectors. This sentence is repeated "
            "to keep native page text above the visual fallback threshold."
        )
        page.insert_textbox(fitz.Rect(50, 50, 590, 310), text, fontsize=12)
    else:
        page.insert_image(page.rect, stream=_png_bytes(kind))
        if kind == "chart":
            page.insert_text((50, 345), "M0_PDF_CHART_QPS_42", fontsize=10)
    document.set_metadata(
        {
            "title": f"M0 synthetic {kind} fixture",
            "author": "RAG-Note M0",
            "creator": "RAG-Note fixture generator",
            "producer": "PyMuPDF",
            "creationDate": "D:20200102030405Z",
            "modDate": "D:20200102030405Z",
        }
    )
    content = document.tobytes(garbage=4, deflate=True, no_new_id=True)
    document.close()
    return content


def _normalize_office_zip(content: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            target.writestr(info, source.read(name))
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = WordDocument()
    properties = document.core_properties
    properties.author = "RAG-Note M0"
    properties.last_modified_by = "RAG-Note M0"
    properties.title = "M0 synthetic DOCX fixture"
    properties.subject = "Synthetic ingestion test data"
    properties.comments = "Contains no user data."
    properties.keywords = "synthetic,m0,ingestion"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    document.add_heading("M0 DOCX ingestion fixture", level=1)
    document.add_paragraph("M0_DOCX_RETRY_TABLE verifies container-aware parsing.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "status"
    table.cell(0, 1).text = "action"
    table.cell(1, 0).text = "429"
    table.cell(1, 1).text = "bounded retry"
    output = io.BytesIO()
    document.save(output)
    return _normalize_office_zip(output.getvalue())


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    properties = presentation.core_properties
    properties.author = "RAG-Note M0"
    properties.last_modified_by = "RAG-Note M0"
    properties.title = "M0 synthetic PPTX fixture"
    properties.subject = "Synthetic ingestion test data"
    properties.comments = "Contains no user data."
    properties.keywords = "synthetic,m0,ingestion"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title_slide.shapes.title.text = "M0 PPTX ingestion fixture"
    title_slide.placeholders[1].text = "M0_PPTX_RETRY_FLOW"

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Retry state"
    table = table_slide.shapes.add_table(
        2, 2, Inches(1), Inches(1.5), Inches(7), Inches(2)
    ).table
    table.cell(0, 0).text = "status"
    table.cell(0, 1).text = "action"
    table.cell(1, 0).text = "429"
    table.cell(1, 1).text = "M0_PPTX_BOUNDED_RETRY"

    output = io.BytesIO()
    presentation.save(output)
    return _normalize_office_zip(output.getvalue())


def build_assets() -> dict[str, bytes]:
    return {
        "m0_text.pdf": _pdf_bytes("text"),
        "m0_chart.pdf": _pdf_bytes("chart"),
        "m0_scan.pdf": _pdf_bytes("scan"),
        "m0_corrupt.pdf": b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\n",
        "m0_retry_table.docx": _docx_bytes(),
        "m0_retry_flow.pptx": _pptx_bytes(),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="生成确定性的 M0 多格式合成样本")
    parser.add_argument(
        "--check", action="store_true", help="只校验现有样本，不写入文件"
    )
    args = parser.parse_args()
    assets = build_assets()
    if args.check:
        mismatches = [
            name
            for name, content in assets.items()
            if not (CORPUS_ROOT / name).is_file()
            or (CORPUS_ROOT / name).read_bytes() != content
        ]
        if mismatches:
            raise SystemExit(f"样本与生成器不一致: {', '.join(mismatches)}")
        return

    CORPUS_ROOT.mkdir(parents=True, exist_ok=True)
    for name, content in assets.items():
        (CORPUS_ROOT / name).write_bytes(content)


if __name__ == "__main__":
    main()
